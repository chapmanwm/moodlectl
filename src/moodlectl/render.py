"""HTML -> readable markdown, and file -> text extraction. Uses pandoc/pdftotext when present."""
from __future__ import annotations

import re
import shutil
import subprocess
import warnings
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup

warnings.filterwarnings("ignore")

JUNK_SELECTORS = [
    "script", "style", "noscript", "svg", "img", "input", "button", "select", "textarea", "nav",
    ".sr-only", ".visually-hidden", ".d-none", ".bulkselect", ".activity-icon", ".dropdown-menu", ".dropdown",
    ".activity-actions", ".course-content-footer", ".section-navigation", "[data-region=loading-icon-container]",
    ".completion-dialog", ".automatic-completion-conditions", ".skip-block", ".modal", ".tiles-top-button",
    "#page-footer", ".activity-header .actions", ".availabilityinfo .btn",
]


STRUCTURE_LISTS = "ul.section, ul.topics, ul.weeks, ul.tiles, ul.img-text, ul[data-for=cmlist], ul.course-section-list, ul.single-section"
STRUCTURE_ITEMS = "li.activity, li.section, li.course-section, li[data-for=cmitem], li[data-for=section]"
KEEP_ATTRS = {"a": {"href"}, "td": {"colspan", "rowspan"}, "th": {"colspan", "rowspan"}}


def main_region(html: str) -> BeautifulSoup:
    soup = BeautifulSoup(html, "lxml")
    region = soup.select_one("#region-main") or soup.select_one("[role=main]") or soup.body or soup
    for sel in JUNK_SELECTORS:
        for el in region.select(sel):
            el.decompose()
    return region


def flatten(region) -> str:
    """Strip Moodle's layout markup so pandoc sees plain content blocks."""
    for el in region.select(STRUCTURE_LISTS):
        el.unwrap()
    for el in region.select(STRUCTURE_ITEMS):
        el.name = "div"
    for lst in region.find_all(["ul", "ol"]):
        if any(getattr(ch, "name", None) not in (None, "li") for ch in lst.children):
            lst.unwrap()
    for h in region.find_all(["h2", "h3", "h4"]):
        if h.get_text(strip=True) in ("Section outline", "Course index"):
            h.decompose()
    for a in region.find_all("a"):
        if not a.get_text(strip=True):
            a.decompose()
    for el in region.find_all(True):
        keep = KEEP_ATTRS.get(el.name, set())
        el.attrs = {k: v for k, v in el.attrs.items() if k in keep}
        if el.name == "span":
            el.unwrap()
    return str(region)


def html_to_md(fragment: str) -> str:
    """Convert an HTML fragment to compact markdown. Falls back to plain text."""
    if shutil.which("pandoc"):
        try:
            out = subprocess.run(
                ["pandoc", "-f", "html-native_divs-native_spans", "-t", "gfm-raw_html", "--wrap=none"],
                input=fragment, capture_output=True, text=True, timeout=60,
            ).stdout
            out = re.sub(r"\n{3,}", "\n\n", out)
            out = re.sub(r"[ \t]+\n", "\n", out)
            return out.strip()
        except Exception:
            pass
    return html_to_text(fragment)


BLOCK_TAGS = ["p", "div", "li", "ul", "ol", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "table", "section", "article", "blockquote", "pre", "dd", "dt"]


def html_to_text(fragment: str) -> str:
    soup = BeautifulSoup(fragment, "lxml")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for el in soup.find_all(BLOCK_TAGS):
        if el.name == "li":
            el.insert(0, "- ")
        el.insert_before("\n")
        el.insert_after("\n")
    for el in soup.find_all(["td", "th"]):
        el.append(" | ")
    text = soup.get_text("")
    text = re.sub(r"[ \t\xa0]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def page_md(html: str) -> str:
    return html_to_md(flatten(main_region(html)))


# -- file -> text -----------------------------------------------------------
TEXT_EXT = {"txt", "csv", "m", "py", "c", "h", "cpp", "json", "tex", "md", "v", "vhd", "asm", "s"}
VIDEO_EXT = {"mp4", "mov", "m4v", "avi", "webm", "mkv", "wmv", "mpg", "mpeg"}
AUDIO_EXT = {"mp3", "wav", "m4a", "aac", "ogg"}


def _run(cmd: list[str]) -> str:
    return subprocess.run(cmd, capture_output=True, text=True, timeout=300).stdout


def extract_text(path: Path) -> str | None:
    ext = path.suffix.lower().lstrip(".")
    try:
        if ext == "pdf":
            if shutil.which("pdftotext"):
                return _run(["pdftotext", "-layout", str(path), "-"])
            from pypdf import PdfReader  # pure-Python fallback ([office] extra)

            return "\n".join(f"--- page {i} ---\n{pg.extract_text() or ''}" for i, pg in enumerate(PdfReader(str(path)).pages, 1))
        if ext == "pptx":
            from pptx import Presentation

            out = []
            for i, s in enumerate(Presentation(str(path)).slides, 1):
                t = [sh.text_frame.text for sh in s.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
                notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
                out.append(f"--- slide {i} ---\n" + "\n".join(t) + (f"\n[notes] {notes}" if notes.strip() else ""))
            return "\n".join(out)
        if ext in ("xlsx", "xlsm"):
            import openpyxl

            out = []
            for ws in openpyxl.load_workbook(str(path), data_only=True, read_only=True):
                out.append(f"=== sheet {ws.title} ===")
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None for c in row):
                        out.append(" | ".join("" if c is None else str(c) for c in row))
            return "\n".join(out)
        if ext in ("docx", "odt", "rtf", "html", "htm", "epub") and shutil.which("pandoc"):
            return _run(["pandoc", "-t", "plain", "--wrap=none", str(path)])
        if ext == "docx":
            import docx  # python-docx fallback ([office] extra)

            d = docx.Document(str(path))
            rows = [" | ".join(c.text for c in r.cells) for tbl in d.tables for r in tbl.rows]
            return "\n".join([p.text for p in d.paragraphs] + rows)
        if ext in TEXT_EXT:
            return path.read_text(errors="ignore")
        if ext == "zip":
            with zipfile.ZipFile(path) as z:
                return "ZIP contents:\n" + "\n".join(z.namelist())
        if ext in ("doc", "ppt", "xls"):
            data = path.read_bytes()
            found = [m.decode("ascii") for m in re.findall(rb"[\x20-\x7e\t]{20,}", data)]
            found += [m.decode("utf-16le") for m in re.findall(rb"(?:[\x20-\x7e]\x00){20,}", data)]
            return "\n".join(x for x in found if "xml" not in x.lower())
    except ImportError:  # optional extractor not installed: pip install 'moodlectl[office]'
        return None
    except Exception as exc:  # extraction is best-effort
        return f"[text extraction failed: {exc}]"
    return None


def safe_name(name: str, limit: int = 120) -> str:
    name = re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", " ", name or "").strip().strip(".")
    name = re.sub(r"\s+", " ", name)
    return (name[:limit].rstrip() or "untitled")
